"""
Generador PD Suficiente / BOA (Vía A) — usa plantilla DOCX con docxtpl.

Plantilla: backend/templates/modelo_pd_fp=.docx
Estructura: 14 secciones A-N (modelo BOA/Aragón)

Uso:
    from generador_pd_suficiente_tpl import generate
    generate(data, out_docx, out_pdf)
"""

import os
from docxtpl import DocxTemplate
from helpers_catalogo import (
    build_ra_desc_map, build_ud_desc_map, build_ce_desc_map,
    resolve_ra_desc, resolve_ud_desc, resolve_ce_desc, resolve_recursos,
    resolve_feoe_dual, _norm_id,
)
from helpers_pd_tablas import insertar_tabla_instrumentos, insertar_tabla_planificacion

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), 'templates', 'modelo_pd_fp=.docx')

DEFAULT_HERRAMIENTAS = [
    "Pizarra y proyector.",
    "Ordenadores e impresora para realización de las prácticas y/o sus informes.",
    "Software ofimático y aplicaciones del módulo.",
    "Herramientas y equipos de medida del laboratorio.",
]

DEFAULT_TEXTO_COORDINACION = (
    "Debido a la dependencia que existe entre los módulos y su profesorado, se ha optado "
    "por establecer reuniones periódicas de coordinación del equipo docente."
)

# Catálogos de checklist codificados en el frontend (METODOLOGIAS de
# MetodologiaTab.tsx, INCLUSION de DiversidadTab.tsx, etc.) -- el backend no
# tiene acceso a los textos traducidos de i18n (ni falta que hace: los
# documentos PD siempre se generan en español, mismo criterio que el resto
# de textos de este generador), así que se duplican aquí solo en español,
# calcados del `defaultValue` de cada checklist.
METODOLOGIAS_LABELS = {
    "ABP": "Aprendizaje Basado en Proyectos",
    "ABR": "Aprendizaje Basado en Retos",
    "FLIP": "Flipped Classroom (Aula Invertida)",
    "COLAB": "Aprendizaje Cooperativo / Colaborativo",
    "SIM": "Simulación de Entornos Profesionales",
    "CASOS": "Método del Caso",
    "GAMIF": "Gamificación / Aprendizaje Basado en Juegos",
    "ApS": "Aprendizaje-Servicio",
    "DEMO": "Demostración Práctica",
    "MAGIS": "Exposición Didáctica Interactiva apoyada en TIC",
    "ETHAZI": "Ethazi / Aprendizaje Colaborativo basado en Retos (ACbR)",
    "AGIL": "Metodologías Ágiles (Design Thinking, Lean Startup, Scrum)",
    "CONTR": "Contrato de Aprendizaje (Learning Contract)",
    "DEBATE": "Debates y Diálogo Educativo",
    "PARES": "Aprendizaje entre Pares (Peer Teaching)",
    "ESTAC": "Estaciones de Aprendizaje",
}

INCLUSION_LABELS = {
    "NIVEL": "Actividades multinivel",
    "AGRUP": "Agrupamientos flexibles y tutoría",
    "TIEMPO": "Flexibilización en tiempos",
    "MATERIAL": "Adaptación de materiales",
    "ACNS": "ACNS (no significativas)",
    "AMPLIA": "Ampliación (altas capacidades)",
}

MODELO_RECUPERACION_LABELS = {
    "R1": "R1 — Recuperación tras 1ª evaluación",
    "R2": "R2 — Recuperación tras 2ª evaluación",
    "R3": "R3 — Recuperación tras 3ª evaluación / final ordinaria",
    "RF": "RF — Recuperación final (tras impartir todas las UD)",
    "EvFE": "EvFE — Evaluación final extraordinaria (segunda convocatoria)",
}

COMPLEMENTARIAS_LABELS = {
    "COMP-VISITA": "Visita técnica a empresa",
    "COMP-CHARLA": "Charla de expertos",
    "COMP-TALLER": "Taller práctico externo",
}

CONTINGENCIA_LABELS = {
    "CONT-ASINC": "Docencia telemática asíncrona",
    "CONT-SINC": "Docencia telemática síncrona",
    "CONT-AUT": "Dosier de tareas autoguiadas",
}

ESCENARIO_CONTINGENCIA_LABELS = {
    "Ausencia de Profesorado": "Ausencia de profesorado",
    "Ausencia de Alumnado": "Ausencia de alumnado",
    "Interrupción Generalizada": "Interrupción generalizada",
    "Otros": "Otros",
}

DEFAULT_TEXTO_PRINCIPIOS_METODOLOGICOS = (
    "La metodología didáctica de la formación profesional específica promueve la integración "
    "de los contenidos científicos, tecnológicos y organizativos, proporcionando una visión "
    "global y coordinada de los procesos productivos en los que debe intervenir el "
    "profesional correspondiente."
)

DEFAULT_TEXTO_EVALUACION_INICIAL = (
    "Esta evaluación tiene como objetivo identificar el nivel de conocimientos y las "
    "necesidades específicas del alumnado, permitiendo ajustar la programación didáctica y "
    "las estrategias de enseñanza a sus características."
)

DEFAULT_TEXTO_E1_DIFERENCIAS_INDIVIDUALES = (
    "Se identifican y respetan las características y necesidades del alumnado, adaptando "
    "las metodologías, los recursos y las evaluaciones a sus capacidades, asegurando que "
    "todos los estudiantes puedan progresar a su propio ritmo."
)

DEFAULT_TEXTO_PLAN_DESDOBLES = (
    "La aplicación de desdobles no se contempla en este módulo, ya que no es necesario "
    "dividir al grupo debido a su tamaño o características. No obstante, se mantienen las "
    "medidas de atención a la diversidad y adaptación curricular no significativa para "
    "asegurar una enseñanza personalizada."
)

DEFAULT_TEXTO_MEDIDAS_BILINGUE = (
    "Este módulo no se imparte en modalidad bilingüe ni se desarrolla dentro de un proyecto "
    "que requiera medidas complementarias específicas."
)

DEFAULT_TEXTO_SEGUIMIENTO = (
    "El seguimiento de la impartición del módulo se realiza a través de los indicadores de "
    "calidad EQAVET (Planificación, Desarrollo y Resultados), autoevaluados al final del "
    "curso, y de las propuestas de mejora derivadas de ese ciclo de mejora continua (PDCA) "
    "para el curso siguiente."
)


def _ra_id_full(ra: dict) -> str:
    id_str = str(ra.get('id_ra', '')).strip().rstrip('.')
    prefix = "" if id_str.upper().startswith("RA") else "RA"
    return f"{prefix}{id_str}"


def _build_context(data: dict) -> dict:
    """
    Construye el contexto Jinja2 para la PD= (Suficiente/BOA).

    data: {
        "modulo": "Maquinas Electricas",
        "ciclo": "Tecnico IEA",
        "departamento": "Electricidad",
        "curso_academico": "2025/2026",
        "horas_totales": 167,
        "horas_semanales": 5,
        "regimen": "diurno",
        "curso_ciclo": "primero",
        "df_ra": [...],
        "df_ud": [...],
        "config_contexto": {...}
    }
    """
    info_mod = data.get("info_modulo") or {}
    modulo = data.get("modulo", "el módulo profesional")
    codigo = info_mod.get("codigo", "")
    if codigo:
        modulo = f"{codigo} - {modulo}"
    ciclo = data.get("ciclo", "el ciclo formativo")
    departamento = data.get("departamento", "")
    # "Departamento Pedagógico": Rafael, 2026-09-13 -- en vez de un nombre de
    # departamento fijo (bug real detectado: la plantilla original tenía
    # "Departamento de Electrónica" grabado a fuego en varios apartados, mal
    # para cualquier módulo de otra familia), se deriva de la familia
    # profesional real del módulo -- ya resuelta por el catálogo oficial en
    # routers/pdf.py, sin necesitar ningún campo nuevo en la UI.
    familia = data.get("familia") or info_mod.get("familia", "")
    departamento_pedagogico = f"Departamento {familia}" if familia else (departamento or "el departamento didáctico responsable del módulo")
    nivel = data.get("nivel", "")
    # La portada (tabla CURSO ACADÉMICO) del documento real solo lleva el año
    # (p.ej. "2025/2026") -- el curso/nivel ("primero", "Grado Medio") ya se
    # dice en prosa dentro de la Introducción, no hace falta repetirlo aquí
    # concatenado (Rafael, 2026-09-13: quitar esa concatenación).
    curso_academico_raw = data.get("curso_academico", "")
    horas_totales = data.get("horas_totales", "")
    horas_semanales = data.get("horas_semanales", "")
    regimen = data.get("regimen", "diurno")
    curso_ciclo = data.get("curso_ciclo", "primero")
    df_ra = data.get("df_ra", [])
    df_ud = data.get("df_ud", [])
    df_ce = data.get("df_ce", [])
    df_sgmt = data.get("df_sgmt", [])
    config = data.get("config_contexto", {})

    # Mapa id_ud -> evaluación (1/2/3) calculado en Planificación mensual
    # (frontend/src/utils/planningGenerator.ts) a partir de la fecha en la
    # que termina cada UD; no hay ningún otro campo en la app que lo derive.
    ev_map = {str(row.get("id_ud", "")): row.get("ev") for row in df_sgmt}

    # Construir mapas de descripciones desde catálogo (df_ra/df_ud/df_ce ->
    # curriculo_data -> catálogo oficial BOA/BOE resuelto en routers/pdf.py)
    ra_desc_map = build_ra_desc_map(data)
    ud_desc_map = build_ud_desc_map(data)
    ce_desc_map = build_ce_desc_map(data)

    context = {
        "modulo": modulo,
        "ciclo": ciclo,
        "departamento": departamento,
        "curso_academico": curso_academico_raw,
        "centro": info_mod.get("centro", ""),
    }

    # ── SECCIÓN A: INTRODUCCIÓN ────────────────────────────────────────
    texto_intro = config.get("texto_introduccion", "")
    if not texto_intro:
        texto_intro = (
            f"Programación didáctica del módulo profesional {modulo}, "
            f"perteneciente al {curso_ciclo} curso del ciclo formativo de grado "
            f"medio {ciclo}, que se imparte en régimen {regimen}, con una duración "
            f"de {horas_totales} periodos lectivos a lo largo del curso académico, "
            f"a razón de {horas_semanales} períodos lectivos semanales."
        )
        # A1_justificacion/A2_contextualizacion: nombres heredados de una
        # numeración interna anterior (ver config_contexto de los .fpp DEMO),
        # con contenido real de justificación/contextualización que hoy no
        # tiene ningún campo de UI con nombre "actual" -- se añaden tal cual
        # si existen, en vez de perder ese texto.
        extra_intro = " ".join(p for p in [config.get("A1_justificacion", ""), config.get("A2_contextualizacion", "")] if p)
        if extra_intro:
            texto_intro = texto_intro + " " + extra_intro
    context["texto_introduccion"] = texto_intro

    texto_uds = config.get("texto_uds_modulo", "")
    if not texto_uds:
        texto_uds = f"El módulo de {modulo}, comprende las siguientes unidades formativas:"
    context["texto_uds_modulo"] = texto_uds

    # Lista de UDs en formato "UF{código}_{n} descripción" — un único bucle
    # dinámico (sustituye a los antiguos placeholders ud1_item..ud11_item,
    # que la plantilla nunca llegó a usar: los dos sitios donde aparecía
    # esta lista tenían el tag genérico sin numerar "{{ ud_item }}", que no
    # existía en el contexto y por tanto salía siempre en blanco).
    list_uf_items = [
        f"UF{codigo}_{i:02d} {resolve_ud_desc(ud, ud_desc_map)}"
        for i, ud in enumerate(df_ud, 1)
    ]
    context["list_uf_items"] = list_uf_items

    # ── SECCIÓN B: FEOE ────────────────────────────────────────────────
    # Ítem 47 (00 IDEAS.md, decisión Rafael 2026-09-22): la unidad real de
    # dualización es el CE (is_dual), no el RA -- un RA puede tener solo
    # parte de sus CE marcados. Dos ramas: para los RA con CE dualizados,
    # se listan los CE concretos que evalúa el tutor de empresa (Anexo XI
    # b); para el resto del módulo, se usa el texto de trabajos
    # alternativos si el profesor lo ha rellenado (Plan FEOE).
    texto_feoe = config.get("texto_feoe", "")
    if not texto_feoe:
        feoe = resolve_feoe_dual(data)
        trabajos_alt = config.get("texto_feoe_trabajos_alternativos", "")
        if feoe["dualizado"]:
            ce_by_ra = {}
            for ce in feoe["ce_dual"]:
                ce_by_ra.setdefault(_norm_id(str(ce.get("id_ra", ""))), []).append(ce)
            partes_dual = []
            for ra in df_ra:
                ces_ra = ce_by_ra.get(_norm_id(str(ra.get("id_ra", ""))))
                if not ces_ra:
                    continue
                ce_ids = ", ".join(str(ce.get("id_ce", "")) for ce in ces_ra)
                partes_dual.append(f"{_ra_id_full(ra)} ({ce_ids})")
            texto_feoe = (
                "En el periodo de Formación en Empresas u Organismos Equiparados (FEOE), la "
                "consecución de los siguientes RA se evalúa a partir de la valoración del tutor/a "
                f"de la empresa u organismo equiparado, según el Anexo XI b: {'; '.join(partes_dual)}."
            )
            horas_imputadas = info_mod.get("horas_imputadas_feoe")
            if horas_imputadas:
                texto_feoe += f" Este módulo tiene imputadas {int(horas_imputadas)} horas lectivas al periodo FEOE."
            ra_no_dual = [ra for ra in df_ra if _norm_id(str(ra.get("id_ra", ""))) not in feoe["ra_ids_dual"]]
            if ra_no_dual and trabajos_alt:
                ids_no_dual = ", ".join(_ra_id_full(ra) for ra in ra_no_dual)
                texto_feoe += f" Para el resto del módulo ({ids_no_dual}), durante ese mismo periodo: {trabajos_alt}"
        else:
            texto_feoe = trabajos_alt or "Este módulo no tiene ningún RA/CE dualizado; no se imparte en modalidad FEOE."
    # B3_vinculacion_empresa: mismo caso que A1/A2 -- texto real bajo un
    # nombre de campo heredado, se añade si existe.
    b3_vinculacion = config.get("B3_vinculacion_empresa", "")
    if b3_vinculacion:
        texto_feoe = texto_feoe + " " + b3_vinculacion
    context["texto_feoe"] = texto_feoe

    # ── SECCIÓN C: SECUENCIACIÓN — Tabla RA×UD (construida con python-docx
    # en generate(), no vía Jinja: así admite tantas filas de UD y columnas
    # de RA como el módulo tenga realmente, en vez del máximo fijo de la
    # plantilla original (11 UD × 7 RA) ─────────────────────────────────
    ra_ids = [_ra_id_full(ra) for ra in df_ra]
    context["tabla_secuenciacion"] = {
        "ra_ids": ra_ids,
        "ra_pcts": [str(int(ra.get('peso_ra', 0) or 0)) for ra in df_ra],
        "filas_ud": [
            {
                "num": str(i),
                "ev": f"{ev_map.get(str(ud.get('id_ud', '')))}ª" if ev_map.get(str(ud.get('id_ud', ''))) else "",
                "nombre": resolve_ud_desc(ud, ud_desc_map),
                "horas": str(ud.get('horas_ud', '')),
                "ra_vals": [
                    (str(int(ud.get(rid, 0))) if ud.get(rid, 0) else "")
                    for rid in ra_ids
                ],
            }
            for i, ud in enumerate(df_ud, 1)
        ],
    }

    # ── SECCIÓN C1: CONTENIDOS — reutiliza el mismo bucle dinámico
    # {%p for ud in list_uds %} que ya trae la plantilla, con el formato
    # "UD n. descripción" que espera esta sección ────────────────────────
    context["list_uds"] = [f"UD {i}. {resolve_ud_desc(ud, ud_desc_map)}" for i, ud in enumerate(df_ud, 1)]

    # ── SECCIÓN C2: CRITERIOS DE EVALUACIÓN — RA + sus CE, número variable
    # de RA y de CE por RA. Se aplana en una única lista de líneas (cada una
    # su propio párrafo vía {%p for %}) en vez de anidar bucles Jinja. ─────
    ce_by_ra: dict = {}
    for ce in df_ce:
        ce_by_ra.setdefault(str(ce.get("id_ra", "")), []).append(ce)

    # Catálogo oficial agrupado por RA (curriculo_data, resuelto en
    # routers/pdf.py vía fetch_curriculo_from_db) — hace falta como
    # respaldo POSICIONAL, no solo por id: el df_ce local suele numerar los
    # CE con letras ("CE1.a", "CE1.b"...) y el catálogo con números
    # ("CE1.1", "CE1.2"...), así que aunque sean los mismos criterios el id
    # normalizado nunca coincide entre uno y otro.
    catalog_ce_by_ra: dict = {}
    for ra_cat in (data.get("curriculo_data") or {}).get("ra") or []:
        rid_cat = str(ra_cat.get("id", "")).strip().rstrip(".")
        catalog_ce_by_ra[rid_cat] = ra_cat.get("ce") or []

    list_c2 = []
    for ra in df_ra:
        rid = str(ra.get("id_ra", ""))
        list_c2.append(f"{_ra_id_full(ra)}. {resolve_ra_desc(ra, ra_desc_map)}")
        list_c2.append("Criterios de evaluación:")

        ces_locales = ce_by_ra.get(rid, [])
        ces_catalogo = catalog_ce_by_ra.get(rid, [])
        mismo_recuento = len(ces_locales) == len(ces_catalogo) and len(ces_catalogo) > 0

        for i, ce in enumerate(ces_locales):
            desc = resolve_ce_desc(ce, ce_desc_map)
            if not desc and mismo_recuento:
                desc = ces_catalogo[i].get("descripcion", "")
            if desc:
                peso = ce.get("peso_ce")
                prefijo = f"({int(peso)}%) " if peso else ""
                list_c2.append(f"{prefijo}{desc}")
    context["list_c2"] = list_c2

    # ── SECCIÓN C3: CRITERIOS DE CALIFICACIÓN ──────────────────────────
    texto_calif = config.get("texto_criterios_calificacion", "")
    if not texto_calif:
        texto_calif = "30%. Conceptuales. 30%. Procedimentales. 40%. Actitudinales."
    context["texto_criterios_calificacion"] = texto_calif

    # Misma tabla "Instrumento | 1er/2º/3er Tri." que PD- (helpers_pd_tablas)
    instrumentos_pct = data.get("instrumentos_pct_trimestre") or []
    if not instrumentos_pct:
        instrumentos_pct = [
            {"nombre": "Exámenes teóricos", "pct_1t": 30, "pct_2t": 20, "pct_3t": 10},
            {"nombre": "Exámenes prácticos", "pct_1t": 20, "pct_2t": 20, "pct_3t": 10},
            {"nombre": "Exposición y defensa proyecto", "pct_1t": 10, "pct_2t": 20, "pct_3t": 30},
            {"nombre": "Informes de ejercicios", "pct_1t": 20, "pct_2t": 30, "pct_3t": 40},
            {"nombre": "Cuaderno de tareas", "pct_1t": 20, "pct_2t": 10, "pct_3t": 10},
        ]
    context["list_instrumentos"] = [
        {
            "nombre": row.get("nombre", ""),
            "pct_1t": f"{row.get('pct_1t', 0)}%",
            "pct_2t": f"{row.get('pct_2t', 0)}%",
            "pct_3t": f"{row.get('pct_3t', 0)}%",
        }
        for row in instrumentos_pct
    ]
    context["pond_1t"] = info_mod.get("pond_1t", 30)
    context["pond_2t"] = info_mod.get("pond_2t", 30)
    context["pond_3t"] = info_mod.get("pond_3t", 40)

    # ── SECCIÓN D: PRINCIPIOS METODOLÓGICOS ────────────────────────────
    # Se aplana en list_d (mismo patrón que list_c2): metodologías activas
    # marcadas + los 3 campos de texto libre relacionados, cada uno su
    # propio párrafo, solo si tienen contenido.
    metodologias_sel = data.get("metodologias_seleccionadas") or []
    list_d = []
    if metodologias_sel:
        nombres = [METODOLOGIAS_LABELS.get(m, m) for m in metodologias_sel]
        list_d.append("Metodologías activas empleadas en el módulo: " + ", ".join(nombres) + ".")
    principios = config.get("principios_metodologicos", "")
    list_d.append(principios if principios else DEFAULT_TEXTO_PRINCIPIOS_METODOLOGICOS)
    estrategias = config.get("estrategias_metodologicas") or config.get("D2_actividades_ea", "")
    if estrategias:
        list_d.append(estrategias)
    colaborativo = config.get("aprendizaje_colaborativo", "")
    if colaborativo:
        list_d.append(colaborativo)
    texto_metodologia_libre = data.get("texto_metodologia_libre", "")
    if texto_metodologia_libre:
        list_d.append(texto_metodologia_libre)
    context["list_d"] = list_d

    # ── SECCIÓN E: EVALUACIÓN INICIAL ───────────────────────────────────
    context["texto_evaluacion_inicial"] = (
        config.get("texto_evaluacion_inicial") or DEFAULT_TEXTO_EVALUACION_INICIAL
    )

    # ── SECCIÓN E1: ATENCIÓN A LAS DIFERENCIAS INDIVIDUALES — checklist +
    # texto libre + los 2 checks booleanos de adaptaciones/flexibilización +
    # tabla DUA (barrera → medida), aplanado en list_e1 ─────────────────
    medidas_inclusion = data.get("medidas_inclusion") or []
    list_e1 = []
    if medidas_inclusion:
        nombres = [INCLUSION_LABELS.get(m, m) for m in medidas_inclusion]
        list_e1.append("Medidas de inclusión aplicadas: " + ", ".join(nombres) + ".")
    f1_diversidad = config.get("F1_diversidad", "")
    list_e1.append(f1_diversidad if f1_diversidad else DEFAULT_TEXTO_E1_DIFERENCIAS_INDIVIDUALES)
    if config.get("adaptaciones_no_significativas"):
        list_e1.append(
            "Se contemplan adaptaciones curriculares no significativas: ajustes metodológicos, "
            "organizativos o de acceso que no alteran los RA ni CE esenciales."
        )
    if config.get("medidas_flexibilizacion"):
        list_e1.append(
            "Se contemplan medidas de flexibilización: alternativas metodológicas en enseñanza "
            "y evaluación que no minorarán las calificaciones."
        )
    texto_inclusion_libre = data.get("texto_inclusion_libre", "")
    if texto_inclusion_libre:
        list_e1.append(texto_inclusion_libre)
    df_dua = data.get("df_dua") or []
    for row in df_dua:
        alumno = row.get("Alumnado_Aula", "")
        barrera = row.get("Barrera", "")
        medida = row.get("Medida_Metodologica") or row.get("Medida_Acceso") or row.get("Medida_Evaluacion", "")
        if alumno or barrera or medida:
            partes = [p for p in [alumno, barrera, medida] if p]
            list_e1.append("DUA — " + " · ".join(partes) + ".")
    context["list_e1"] = list_e1

    # ── SECCIÓN F: PROCEDIMIENTOS E INSTRUMENTOS DE EVALUACIÓN ──────────
    texto_eval_info = data.get("textos_pd_eval_informacion", "")
    context["texto_procedimientos_instrumentos"] = texto_eval_info or (
        "Los procedimientos e instrumentos de evaluación se ajustan a los objetivos y "
        "competencias del módulo, garantizando una evaluación formativa y final que facilite "
        "el seguimiento y la mejora del proceso educativo."
    )

    # ── SECCIÓN G: ACTIVIDADES DE RECUPERACIÓN Y REFUERZO — checklist de
    # convocatorias (R1/R2/R3/RF/EvFE) + textos libres, aplanado en list_g ─
    modelo_recup = config.get("modelo_recuperacion") or []
    list_g = []
    if modelo_recup:
        nombres = [MODELO_RECUPERACION_LABELS.get(m, m) for m in modelo_recup]
        list_g.append("Modelo de recuperación aplicado: " + ", ".join(nombres) + ".")
    texto_recuperacion = (
        data.get("textos_pd_eval_recuperacion", "")
        or config.get("texto_modelo_recuperacion", "")
        or config.get("E5_recuperacion", "")  # nombre heredado, ver nota de A1/A2 arriba
    )
    list_g.append(texto_recuperacion if texto_recuperacion else (
        "Las actividades de recuperación y refuerzo se fundamentan en los resultados de "
        "aprendizaje en los que el alumnado haya tenido dificultades, e incluyen actividades "
        "de ampliación para el alumnado que las precise."
    ))
    texto_perdida_continua = data.get("textos_pd_eval_perdida_continua", "")
    if texto_perdida_continua:
        list_g.append(texto_perdida_continua)
    context["list_g"] = list_g

    # ── SECCIÓN G1: PLAN DE RECUPERACIÓN — convocatoria extraordinaria y
    # alumnado con el módulo pendiente ──────────────────────────────────
    texto_pendientes = data.get("textos_pd_eval_pendientes", "")
    # p_ev: % legal de Pérdida de Evaluación Continua (PdEvC), campo real y
    # editable por el docente en Programación > Módulo > Datos (por defecto
    # 15%, ver AttendanceAccumulated.tsx) -- no un valor fijo del 15%.
    p_ev = info_mod.get("p_ev") or 15
    context["texto_plan_recuperacion"] = texto_pendientes or (
        f"El alumnado que, no superando el {p_ev}% de faltas de evaluación (umbral de "
        "Pérdida de Evaluación Continua, PdEvC), haya suspendido el módulo en convocatoria "
        "ordinaria, podrá recuperarlo en convocatoria extraordinaria, aplicándose en "
        "cualquier caso los criterios de calificación detallados en esta programación."
    )

    # ── SECCIÓN H: RESULTADOS DE APRENDIZAJE — número variable de RA ────
    context["n_ra"] = len(df_ra)
    context["list_ras_h"] = [f"RA {i}. {resolve_ra_desc(ra, ra_desc_map)}" for i, ra in enumerate(df_ra, 1)]

    # ── SECCIÓN I: PLAN DE APLICACIÓN DE LOS DESDOBLES ─────────────────
    context["texto_plan_desdobles"] = (
        config.get("plan_desdobles") or config.get("D3_agrupamientos") or DEFAULT_TEXTO_PLAN_DESDOBLES
    )

    # ── SECCIÓN J: HERRAMIENTAS — resuelve recursos_espacios (catálogo
    # codificado en helpers_catalogo.RECURSOS_LABELS + añadidos libres del
    # profesorado), más G1_infraestructuras/G2_herramientas_tic (nombres
    # heredados, texto libre real) si existen ──────────────────────────
    recursos = data.get("recursos_espacios") or []
    list_herramientas = resolve_recursos(recursos) if recursos else list(DEFAULT_HERRAMIENTAS)
    g1_infra = config.get("G1_infraestructuras", "")
    if g1_infra:
        list_herramientas.append(g1_infra)
    g2_tic = config.get("G2_herramientas_tic", "")
    if g2_tic:
        list_herramientas.append(g2_tic)
    context["list_herramientas"] = list_herramientas
    context["textos_pd_bibliografia"] = (
        data.get("textos_pd_bibliografia") or config.get("G3_bibliografia") or ""
    )

    # ── SECCIÓN K: ACTIVIDADES COMPLEMENTARIAS Y EXTRAESCOLARES — solo
    # complementarias (en horario lectivo, evaluables); las extraescolares
    # no forman parte de la programación didáctica (van en la PGA), mismo
    # criterio ya documentado en InnovacionTab.tsx ─────────────────────
    actividades_compl = data.get("actividades_complementarias") or []
    list_k = []
    if actividades_compl:
        nombres = [COMPLEMENTARIAS_LABELS.get(a, a) for a in actividades_compl]
        list_k.append("Actividades complementarias previstas: " + ", ".join(nombres) + ".")
    h1_complementarias = config.get("H1_complementarias", "")
    if h1_complementarias:
        list_k.append(h1_complementarias)
    if not list_k:
        list_k.append(
            f"Dado que existe una estrecha interdependencia entre los módulos y el "
            f"profesorado del {departamento_pedagogico}, las actividades complementarias "
            "se desarrollarán de manera coordinada con el resto del equipo docente."
        )
    context["list_k"] = list_k

    # ── SECCIÓN L: MEDIDAS COMPLEMENTARIAS EN PROYECTOS O BILINGÜES ────
    context["texto_medidas_bilingue"] = (
        config.get("texto_medidas_bilingue") or DEFAULT_TEXTO_MEDIDAS_BILINGUE
    )

    # ── SECCIÓN M: MECANISMOS DE SEGUIMIENTO Y VALORACIÓN — indicadores
    # EQAVET (Planificación/Desarrollo/Resultados) ya puntuados + propuestas
    # de mejora PDCA, si existen (normalmente se rellenan a fin de curso;
    # con el módulo recién empezado se usa el texto por defecto) ───────
    eqavet = data.get("eqavet_evaluacion") or {}
    list_m = []
    categorias_eqavet = [("planificacion", "Planificación"), ("desarrollo", "Desarrollo"), ("resultados", "Resultados")]
    for key, label in categorias_eqavet:
        fuertes = eqavet.get(f"puntos_fuertes_{key}", "")
        mejora = eqavet.get(f"areas_mejora_{key}", "")
        if fuertes:
            list_m.append(f"{label} — puntos fuertes: {fuertes}")
        if mejora:
            list_m.append(f"{label} — áreas de mejora: {mejora}")
    if not list_m:
        list_m.append(
            f"Periódicamente el {departamento_pedagogico} realiza la revisión del "
            "cumplimiento de esta programación didáctica, así como la evolución de los "
            "resultados académicos del alumnado, analizando la situación y proponiendo "
            "acciones de mejora."
        )
        list_m.append(DEFAULT_TEXTO_SEGUIMIENTO)
    context["list_m"] = list_m

    # ── SECCIÓN N: PLAN DE CONTINGENCIA — coordinación con otros módulos
    # (textos_pd_metodologia_labor_coordinada) + los 3 escenarios reales
    # (profesor/alumnado/interrupción general) + checklist de medidas +
    # registro de escenarios (df_contingencia), aplanado en list_n ─────
    list_n = [
        data.get("textos_pd_metodologia_labor_coordinada") or DEFAULT_TEXTO_COORDINACION
    ]
    medidas_contingencia = data.get("medidas_contingencia") or []
    if medidas_contingencia:
        nombres = [CONTINGENCIA_LABELS.get(m, m) for m in medidas_contingencia]
        list_n.append("Medidas de contingencia previstas: " + ", ".join(nombres) + ".")
    contingencia_profesor = config.get("contingencia_profesor", "")
    if contingencia_profesor:
        list_n.append("Ausencia del profesorado: " + contingencia_profesor)
    else:
        list_n.append(
            "Ante la ausencia del profesorado titular, se facilitarán las tareas o unidades a "
            "estudiar por los canales de comunicación habituales del centro. En caso de baja de "
            "larga duración, las actividades docentes corresponderán al nuevo docente o, si la "
            f"baja no es cubierta, al {departamento_pedagogico}."
        )
    contingencia_alumnado = config.get("contingencia_alumnado", "")
    if contingencia_alumnado:
        list_n.append("Ausencia del alumnado: " + contingencia_alumnado)
    contingencia_general = config.get("contingencia_general") or config.get("J3_contingencia", "")
    if contingencia_general:
        list_n.append("Interrupción generalizada de las clases: " + contingencia_general)
    texto_contingencia_libre = data.get("texto_contingencia_libre", "")
    if texto_contingencia_libre:
        list_n.append(texto_contingencia_libre)
    df_contingencia = data.get("df_contingencia") or []
    for row in df_contingencia:
        escenario = ESCENARIO_CONTINGENCIA_LABELS.get(row.get("Escenario", ""), row.get("Escenario", ""))
        actividades = row.get("Actividades", "")
        if escenario and actividades:
            list_n.append(f"{escenario}: {actividades}")
    context["list_n"] = list_n
    # Se mantiene por compatibilidad con la plantilla actual (tag suelto);
    # el nuevo bucle {%p for line in list_n %} lo sustituye por completo.
    context["textos_pd_metodologia_labor_coordinada"] = (
        data.get("textos_pd_metodologia_labor_coordinada") or DEFAULT_TEXTO_COORDINACION
    )

    return context


def generate(data: dict, out_docx: str, out_pdf: str = None):
    """
    Genera la PD= (Suficiente/BOA) usando la plantilla DOCX.
    """
    if not os.path.exists(TEMPLATE_PATH):
        raise FileNotFoundError(
            f"No se encontró la plantilla en: {TEMPLATE_PATH}. "
            f"Ejecuta 'python scripts/preparar_plantilla_pd_suficiente.py' para generarla."
        )

    tpl = DocxTemplate(TEMPLATE_PATH)
    context = _build_context(data)
    # autoescape=True: ver nota en generador_pd_jeg.py -- sin esto, un '<' o
    # '&' suelto en texto libre del profesor se inserta tal cual en el XML
    # del docx y puede dejar el documento mal formado.
    tpl.render(context, autoescape=True)

    doc = tpl.docx
    _insertar_tabla_secuenciacion(doc, context["tabla_secuenciacion"])
    for p in doc.paragraphs:
        if "[[TABLA_INSTRUMENTOS_C3]]" in p.text:
            p.text = p.text.replace("[[TABLA_INSTRUMENTOS_C3]]", "")
            insertar_tabla_instrumentos(
                doc, p, context["list_instrumentos"],
                context["pond_1t"], context["pond_2t"], context["pond_3t"],
            )
            break

    # ── Página final: previsión de planificación mensual (item 9 backlog) ──
    insertar_tabla_planificacion(
        doc, data.get("df_ud", []),
        data.get("info_fechas", {}), data.get("horario", {}), data.get("calendar_notes", {}),
    )

    tpl.save(out_docx)


def _insertar_tabla_secuenciacion(doc, tabla_data: dict):
    """
    Construye con python-docx la tabla de la sección C (Secuenciación),
    en el marcador [[TABLA_SECUENCIACION]]: N | Ev. | Unidades didácticas |
    H. | RA1..RAn (tantas columnas de RA como tenga el módulo) y una fila
    de "% por RA" además de una fila por cada UD.
    """
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from helpers_pd_tablas import aplicar_bordes_rejilla

    ra_ids = tabla_data["ra_ids"]
    ra_pcts = tabla_data["ra_pcts"]
    filas_ud = tabla_data["filas_ud"]
    n_cols = 4 + len(ra_ids)

    for p in doc.paragraphs:
        if "[[TABLA_SECUENCIACION]]" not in p.text:
            continue
        p.text = p.text.replace("[[TABLA_SECUENCIACION]]", "")

        table = doc.add_table(rows=2 + len(filas_ud), cols=n_cols)
        aplicar_bordes_rejilla(table)

        # Anchos: 18cm totales repartidos entre las 4 columnas fijas (más
        # anchas) y las columnas de RA (más estrechas), igual de estrecho
        # cuantas más columnas de RA haya para no desbordar la página.
        fixed_widths = [Cm(1.0), Cm(1.2), Cm(7.0), Cm(1.2)]
        resto = 18.0 - sum(w.cm for w in fixed_widths)
        ra_w = Cm(max(resto / len(ra_ids), 0.8)) if ra_ids else Cm(0)
        col_widths = fixed_widths + [ra_w] * len(ra_ids)
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = w

        def _fmt_cell(cell, text, bold=False, center=True):
            cell.text = str(text)
            for para in cell.paragraphs:
                if center:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:
                    r.bold = bold
                    r.font.name = 'Arial'
                    r.font.size = Pt(8)
                    r.font.color.rgb = RGBColor(0, 0, 0)

        # Fila 0: "% por RA" en las columnas de RA
        pct_cells = table.rows[0].cells
        _fmt_cell(pct_cells[2], "% por RA", bold=True, center=False)
        for j, pct in enumerate(ra_pcts):
            _fmt_cell(pct_cells[4 + j], f"{pct}%", bold=True)

        # Fila 1: cabeceras de columna
        hdr_cells = table.rows[1].cells
        for i, label in enumerate(["N", "Ev.", "Unidades didácticas", "H."]):
            _fmt_cell(hdr_cells[i], label, bold=True, center=(i != 2))
        for j, rid in enumerate(ra_ids):
            _fmt_cell(hdr_cells[4 + j], rid, bold=True)

        # Filas de datos, una por UD
        for i, fila in enumerate(filas_ud):
            row_cells = table.rows[2 + i].cells
            _fmt_cell(row_cells[0], fila["num"])
            _fmt_cell(row_cells[1], fila["ev"])
            _fmt_cell(row_cells[2], fila["nombre"], center=False)
            _fmt_cell(row_cells[3], fila["horas"])
            for j, val in enumerate(fila["ra_vals"]):
                _fmt_cell(row_cells[4 + j], val)

        p._p.addnext(table._tbl)
        p._p.getparent().remove(p._p)
        break
