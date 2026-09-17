# -*- coding: utf-8 -*-
"""
pdf_cobertura_ce.py
PDF A4 apaisado - Matriz de cobertura CE x Instrumento: que instrumentos
evaluan cada Criterio de Evaluacion, a traves de la cadena real
Instrumento -> Indicador -> CE (misma logica que la matriz interactiva de
Instrumentos > Modelo JEG, JegModeloTab.tsx). Adoptado de
ResumenCriteriosPorTarea.js (app de referencia de un companero, David),
coloreado por evaluacion (Ev1/Ev2/Ev3) con la misma paleta morado/verde
azulado/ambar que usa el calendario para los trimestres -- Rafael, 2026-09-17.
"""
import io
import pandas as pd
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import (
    BaseDocTemplate, Frame, PageTemplate,
    Table, TableStyle, Paragraph, Spacer
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib.enums import TA_CENTER, TA_LEFT

# Misma familia de color que la matriz interactiva (JegModeloTab.tsx,
# COLOR_POR_EVALUACION: purple/teal/amber) y que el calendario usa para los
# trimestres -- aqui como tinta legible sobre fondo blanco (impreso), no
# como superposicion translucida sobre fondo oscuro (pantalla).
COLOR_POR_EVALUACION = {
    "Ev1": {"header": "#7e22ce", "bg": "#f3e9fe"},
    "Ev2": {"header": "#0f766e", "bg": "#e3f7f5"},
    "Ev3": {"header": "#b45309", "bg": "#fdf1dc"},
    "EvFO": {"header": "#666666", "bg": "#eeeeee"},
    "EvFE": {"header": "#666666", "bg": "#eeeeee"},
}
_COLOR_DEFECTO = {"header": "#666666", "bg": "#eeeeee"}


def _draw_page_decorations(canv, doc):
    canv.saveState()
    W, H = landscape(A4)
    canv.setFont("Helvetica-Bold", 10)
    canv.setFillColor(colors.HexColor("#777777"))
    canv.drawCentredString(W / 2, H - 1.5 * cm, doc.cal_titulo)
    canv.setFont("Helvetica", 9)
    canv.drawRightString(W - 1 * cm, 1 * cm, doc.cal_pie)
    canv.restoreState()


def _cobertura_ce(df_ce: pd.DataFrame, df_indicadores: pd.DataFrame, df_instr: pd.DataFrame):
    """Para cada CE, la lista de id_instrumento que lo cubren -- replica
    exactamente indsDeEsteCe/instrumentosQueCubren de JegModeloTab.tsx."""
    filas = []
    if df_ce.empty or "id_ce" not in df_ce.columns:
        return filas
    for _, ce_row in df_ce.iterrows():
        id_ce = ce_row.get("id_ce")
        if not id_ce:
            continue
        if not df_indicadores.empty and "id_ce" in df_indicadores.columns:
            inds_de_este_ce = df_indicadores.loc[df_indicadores["id_ce"] == id_ce, "id_indicador"].tolist()
        else:
            inds_de_este_ce = []
        cubren = []
        for _, instr_row in df_instr.iterrows():
            vinculados = instr_row.get("indicadores_vinculados") or []
            if any(i in inds_de_este_ce for i in vinculados):
                cubren.append(instr_row.get("id_instrumento"))
        filas.append({"id_ce": id_ce, "cubren": cubren})
    return filas


def generar_pdf_cobertura_ce(
    info_modulo: dict,
    df_ce: pd.DataFrame,
    df_indicadores: pd.DataFrame,
    df_instr: pd.DataFrame,
):
    buffer = io.BytesIO()
    W, H = landscape(A4)
    left_m, right_m, top_m, bottom_m = 2.0 * cm, 1.0 * cm, 2.0 * cm, 1.5 * cm

    doc = BaseDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=left_m, rightMargin=right_m, topMargin=top_m, bottomMargin=bottom_m,
    )
    nombre_modulo = info_modulo.get("modulo", "Módulo")
    doc.cal_titulo = f"Matriz de cobertura CE × Instrumento  ·  {nombre_modulo}"
    doc.cal_pie = f"{info_modulo.get('centro', '')} ({info_modulo.get('profesorado', '')})"

    frame = Frame(left_m, bottom_m, W - left_m - right_m, H - top_m - bottom_m, id="main")
    doc.addPageTemplates([PageTemplate(id="port", frames=[frame], onPage=_draw_page_decorations)])

    styles = getSampleStyleSheet()
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=14, spaceAfter=8, textColor=colors.black, spaceBefore=4)
    norm_left = ParagraphStyle("NorL", parent=styles["Normal"], fontSize=9, leading=11, alignment=TA_LEFT)
    smlB = ParagraphStyle("SmB", parent=styles["Normal"], fontSize=8, leading=9, fontName="Helvetica-Bold", alignment=TA_CENTER)
    leyenda = ParagraphStyle("Leyenda", parent=styles["Normal"], fontSize=8, leading=11, alignment=TA_LEFT, textColor=colors.HexColor("#555555"))

    elements = [Paragraph("Matriz de cobertura CE × Instrumento", h2)]

    if df_ce.empty or df_instr.empty:
        elements.append(Paragraph("<i>No hay Criterios de evaluación o Instrumentos definidos.</i>", norm_left))
        doc.build(elements)
        buffer.seek(0)
        return buffer

    leyenda_txt = " &nbsp;&nbsp; ".join(
        f'<font color="{COLOR_POR_EVALUACION.get(ev, _COLOR_DEFECTO)["header"]}">●</font> {ev}'
        for ev in ["Ev1", "Ev2", "Ev3", "EvFO / EvFE"]
    )
    elements.append(Paragraph(leyenda_txt, leyenda))
    elements.append(Spacer(1, 10))

    filas = _cobertura_ce(df_ce, df_indicadores, df_instr)
    instr_ids = df_instr["id_instrumento"].tolist() if "id_instrumento" in df_instr.columns else []
    evaluacion_por_instr = dict(zip(df_instr.get("id_instrumento", []), df_instr.get("evaluacion", [])))

    header = [Paragraph("<b>CE</b>", smlB)]
    for instr_id in instr_ids:
        color = COLOR_POR_EVALUACION.get(evaluacion_por_instr.get(instr_id), _COLOR_DEFECTO)
        header.append(Paragraph(f'<font color="{color["header"]}"><b>{instr_id}</b></font>', smlB))
    header.append(Paragraph("<b>Cobertura</b>", smlB))

    data = [header]
    row_bg_cmds = []
    for r_idx, fila in enumerate(filas, start=1):
        cubren = fila["cubren"]
        cubierto = len(cubren) > 0
        row = [Paragraph(f"<b>{fila['id_ce']}</b>", smlB)]
        for instr_id in instr_ids:
            if instr_id in cubren:
                color = COLOR_POR_EVALUACION.get(evaluacion_por_instr.get(instr_id), _COLOR_DEFECTO)
                row.append(Paragraph(f'<font color="{color["header"]}"><b>✓</b></font>', smlB))
                col_idx = 1 + instr_ids.index(instr_id)
                row_bg_cmds.append(("BACKGROUND", (col_idx, r_idx), (col_idx, r_idx), colors.HexColor(color["bg"])))
            else:
                row.append("")
        if cubierto:
            row.append(Paragraph(f"<font color='#228B22'>{len(cubren)} instr.</font>", smlB))
        else:
            row.append(Paragraph("<font color='#cc0000'><b>Sin cobertura</b></font>", smlB))
            row_bg_cmds.append(("BACKGROUND", (0, r_idx), (-1, r_idx), colors.HexColor("#fdecea")))
        data.append(row)

    avail_w = W - left_m - right_m
    ce_w, cob_w = 2.2 * cm, 3 * cm
    n_instr = len(instr_ids)
    instr_w = max(1.4 * cm, (avail_w - ce_w - cob_w) / max(n_instr, 1))
    col_widths = [ce_w] + [instr_w] * n_instr + [cob_w]

    table = Table(data, colWidths=col_widths, repeatRows=1)
    ts_cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0f0f0")),
        ("LINEBELOW", (0, 0), (-1, 0), 1.5, colors.HexColor("#222222")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BOX", (0, 0), (-1, -1), 1.5, colors.HexColor("#222222")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#bbbbbb")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ] + row_bg_cmds
    table.setStyle(TableStyle(ts_cmds))
    elements.append(table)

    doc.build(elements)
    buffer.seek(0)
    return buffer


def generar_docx_cobertura_ce(
    info_modulo: dict,
    df_ce: pd.DataFrame,
    df_indicadores: pd.DataFrame,
    df_instr: pd.DataFrame,
):
    """Versión .docx editable de la matriz de cobertura CE × Instrumento."""
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx_helpers import new_document, add_title, add_meta_line, add_section_heading, _shade_cell, doc_to_bytes

    doc = new_document(landscape=True)
    add_title(doc, "Matriz de cobertura CE × Instrumento", info_modulo.get("modulo", "Módulo"))
    add_meta_line(doc, f"{info_modulo.get('centro', '')} ({info_modulo.get('profesorado', '')})")

    if df_ce.empty or df_instr.empty:
        doc.add_paragraph("No hay Criterios de evaluación o Instrumentos definidos.")
        return doc_to_bytes(doc)

    filas = _cobertura_ce(df_ce, df_indicadores, df_instr)
    instr_ids = df_instr["id_instrumento"].tolist() if "id_instrumento" in df_instr.columns else []
    evaluacion_por_instr = dict(zip(df_instr.get("id_instrumento", []), df_instr.get("evaluacion", [])))

    leyenda_p = doc.add_paragraph()
    for i, ev in enumerate(["Ev1", "Ev2", "Ev3", "EvFO / EvFE"]):
        if i > 0:
            leyenda_p.add_run("   ")
        color = COLOR_POR_EVALUACION.get(ev, _COLOR_DEFECTO)
        run_dot = leyenda_p.add_run("● ")
        run_dot.font.color.rgb = RGBColor.from_string(color["header"].lstrip("#"))
        leyenda_p.add_run(ev)
    for run in leyenda_p.runs:
        run.font.size = Pt(9)

    add_section_heading(doc, "Cobertura")

    headers = ["CE"] + instr_ids + ["Cobertura"]
    table = doc.add_table(rows=1 + len(filas), cols=len(headers))
    table.style = "Table Grid"

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = str(h)
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(9)
        _shade_cell(hdr_cells[i], "E0E0E0")
        if i > 0 and i <= len(instr_ids):
            color = COLOR_POR_EVALUACION.get(evaluacion_por_instr.get(instr_ids[i - 1]), _COLOR_DEFECTO)
            for p in hdr_cells[i].paragraphs:
                for run in p.runs:
                    run.font.color.rgb = RGBColor.from_string(color["header"].lstrip("#"))

    for r_idx, fila in enumerate(filas):
        cells = table.rows[r_idx + 1].cells
        cubren = fila["cubren"]
        cubierto = len(cubren) > 0
        cells[0].text = str(fila["id_ce"])
        for p in cells[0].paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(9)
        for c_idx, instr_id in enumerate(instr_ids, start=1):
            covered = instr_id in cubren
            cells[c_idx].text = "✓" if covered else ""
            for p in cells[c_idx].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.bold = True
            if covered:
                color = COLOR_POR_EVALUACION.get(evaluacion_por_instr.get(instr_id), _COLOR_DEFECTO)
                for p in cells[c_idx].paragraphs:
                    for run in p.runs:
                        run.font.color.rgb = RGBColor.from_string(color["header"].lstrip("#"))
                _shade_cell(cells[c_idx], color["bg"].lstrip("#"))
        cob_cell = cells[len(instr_ids) + 1]
        cob_cell.text = f"{len(cubren)} instr." if cubierto else "Sin cobertura"
        for p in cob_cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor.from_string("228B22" if cubierto else "CC0000")
                if not cubierto:
                    run.font.bold = True
        if not cubierto:
            for c in cells:
                _shade_cell(c, "FDECEA")

    return doc_to_bytes(doc)
