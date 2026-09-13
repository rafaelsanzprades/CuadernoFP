# -*- coding: utf-8 -*-
"""
Arregla la plantilla modelo_pd_jeg_tpl_final.docx (PD+/JEG): sustituye
ocurrencias de {{ campo_libre }} -- y algunos parrafos totalmente vacios --
por tags Jinja especificos que ya construye generador_pd_jeg.py::_build_context
pero que nunca llegaban a la plantilla real.

No se toca la estructura del documento (ni parrafos ni tablas nuevas): solo
se cambia el nombre del tag dentro de parrafos ya existentes, o se rellena un
parrafo vacio con un tag. Los bloques [[ ... ]] (menus de opciones para que
el docente elija/edite a mano) se dejan intactos, salvo por el tag interno
que sustituyen -- no se quita ningun corchete.

Uso: .venv/Scripts/python.exe scripts/arreglar_plantilla_jeg.py
"""
import os
import docx

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "..", "templates", "modelo_pd_jeg_tpl_final.docx")

# indice de parrafo (segun doc.paragraphs) -> nuevo tag Jinja (sin llaves)
REEMPLAZOS_CAMPO_LIBRE = {
    490: "texto_criterios_calificacion",
    597: "actividad1_ra",
    599: "actividad1_entidad",
    607: "actividad2_ra",
    609: "actividad2_entidad",
    631: "recursos_personales",
    632: "espacios_recursos",
    633: "recursos_materiales",
    634: "recursos_documentales",
    635: "recursos_digitales",
    656: "recursos_personales",
    657: "espacios_recursos",
    658: "recursos_materiales",
    659: "recursos_documentales",
    660: "recursos_digitales",
    680: "recursos_personales",
    681: "espacios_recursos",
    682: "recursos_materiales",
    683: "recursos_documentales",
    684: "recursos_digitales",
}

# indice de parrafo totalmente vacio -> tag Jinja a insertar (sin llaves)
RELLENOS_PARRAFO_VACIO = {
    84: "espacios_recursos",
    86: "equipamiento_tecnico",
    88: "recursos_tecnologicos",
    90: "software_herramientas",
    92: "material_didactico",
    94: "espacios_complementarios",
}


def _reemplazar_campo_libre(paragraph, nuevo_tag):
    objetivo = "campo_libre"
    nuevo = nuevo_tag
    for run in paragraph.runs:
        if objetivo in run.text:
            run.text = run.text.replace(objetivo, nuevo, 1)
            return True
    return False


def main():
    doc = docx.Document(TEMPLATE_PATH)
    paras = doc.paragraphs

    fallos = []

    for idx, nuevo_tag in REEMPLAZOS_CAMPO_LIBRE.items():
        p = paras[idx]
        if "campo_libre" not in p.text:
            fallos.append((idx, "no contiene 'campo_libre'", p.text[:80]))
            continue
        if not _reemplazar_campo_libre(p, nuevo_tag):
            fallos.append((idx, "campo_libre partido entre runs", p.text[:80]))

    for idx, tag in RELLENOS_PARRAFO_VACIO.items():
        p = paras[idx]
        if p.text.strip() != "":
            fallos.append((idx, "parrafo no estaba vacio", p.text[:80]))
            continue
        if p.runs:
            p.runs[0].text = "{{ " + tag + " }}"
        else:
            p.add_run("{{ " + tag + " }}")

    if fallos:
        print("FALLOS:")
        for f in fallos:
            print(" ", f)
        raise SystemExit(1)

    doc.save(TEMPLATE_PATH)
    print("OK -- plantilla JEG actualizada:", len(REEMPLAZOS_CAMPO_LIBRE), "tags renombrados,",
          len(RELLENOS_PARRAFO_VACIO), "parrafos vacios rellenados.")


if __name__ == "__main__":
    main()
