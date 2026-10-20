"""
Amplía la plantilla PD+ (modelo_pd_jeg_tpl_final.docx) con datos de la app que antes no llegaban a ningún documento
(2026-10-20, a petición de Rafael). Solo añade; no cambia ni quita nada del texto del modelo:

  1. Tabla de «Identificación»: tres filas nuevas tras «Nombre» del centro: Localidad, Provincia y Comunidad autónoma
     ({{ municipio }}, {{ provincia }}, {{ comunidad_autonoma }}).
  2. Apartado «Actividades, técnicas e instrumentos»: un bucle {%p for x in list_actividades_eval %} que lista, por trimestre,
     las actividades de evaluación de Instrumentos › Trimestres con su peso, justo tras «…tipos de actividades de evaluación:».

Los proyectos y retos (apartado «Aprendizaje colaborativo…») y las sesiones (Contenidos) se rellenan desde el generador, sin tocar la plantilla.

Es idempotente. Uso:   cd backend && python scripts/ampliar_plantilla_jeg_campos.py
"""

import copy
import os

from docx import Document

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RUTA = os.path.join(BASE, "templates", "modelo_pd_jeg_tpl_final.docx")


def poner_texto(parrafo, texto):
    if not parrafo.runs:
        parrafo.add_run(texto)
        return
    parrafo.runs[0].text = texto
    for extra in parrafo.runs[1:]:
        extra.text = ""


def main():
    doc = Document(RUTA)
    cambios = []

    # 1. filas de Identificación
    tabla = doc.tables[0]
    textos = [c.text.strip() for c in tabla.rows[1].cells]
    if not any("{{ municipio }}" in c.text for r in tabla.rows for c in r.cells):
        fila_nombre = tabla.rows[1]
        ancla = fila_nombre._tr
        for etiqueta, tag in (("Comunidad autónoma", "{{ comunidad_autonoma }}"), ("Provincia", "{{ provincia }}"), ("Localidad", "{{ municipio }}")):
            nueva = copy.deepcopy(fila_nombre._tr)
            ancla.addnext(nueva)
            from docx.table import _Row
            fila = _Row(nueva, tabla)
            poner_texto(fila.cells[1].paragraphs[0], etiqueta)
            poner_texto(fila.cells[2].paragraphs[0], tag)
        cambios.append("Identificación: filas Localidad, Provincia y Comunidad autónoma")

    # 2. bucle de actividades de evaluación
    if not any("list_actividades_eval" in p.text for p in doc.paragraphs):
        for p in doc.paragraphs:
            if p.text.strip().startswith("A lo largo del desarrollo del") and "tipos de actividades de evaluación" in p.text:
                fin = copy.deepcopy(p._p)
                cuerpo = copy.deepcopy(p._p)
                ini = copy.deepcopy(p._p)
                p._p.addnext(fin)
                p._p.addnext(cuerpo)
                p._p.addnext(ini)
                from docx.text.paragraph import Paragraph
                poner_texto(Paragraph(ini, p._parent), "{%p for x in list_actividades_eval %}")
                poner_texto(Paragraph(cuerpo, p._parent), "{{ x }}")
                poner_texto(Paragraph(fin, p._parent), "{%p endfor %}")
                cambios.append("4.1: bucle list_actividades_eval")
                break
        else:
            raise SystemExit("No se encontró el párrafo de «tipos de actividades de evaluación»")

    if cambios:
        doc.save(RUTA)
    print("Cambios:" if cambios else "Sin cambios (ya estaba ampliada).")
    for c in cambios:
        print("  -", c)


if __name__ == "__main__":
    main()
