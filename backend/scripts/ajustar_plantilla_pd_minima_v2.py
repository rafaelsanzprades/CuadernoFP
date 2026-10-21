"""PD- (resumen para el alumnado): ajusta templates/modelo_pd_fp-.docx al modelo de Rafael (2026-10-22).

Cambios: título a 12 pt; cabecera en tres líneas con etiquetas en negrita (Módulo / Título, Grupo, Curso /
Centro, Departamento); fuera la frase «Seguro que tienes muchas preguntas…» y el texto final «¡ÁNIMO!»;
encabezados a 9 pt; párrafos de texto con sangría de 0,5 cm. Idempotente.
"""
import os
import sys

from docx import Document
from docx.shared import Cm, Pt

RUTA = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "templates", "modelo_pd_fp-.docx")

doc = Document(RUTA)
pars = doc.paragraphs

for p in list(pars):
    t = p.text.strip()
    if t.startswith("Seguro que tienes muchas preguntas") or t == "{{ texto_final }}":
        p._p.getparent().remove(p._p)

for p in doc.paragraphs:
    t = p.text
    if t.startswith("Resumen de la Programación"):
        for r in p.runs:
            r.font.size = Pt(12)
    elif t.startswith("Módulo:"):
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        trozos = [
            ("Módulo:", True), (" {{ modulo }}", False),
            ("\nTítulo:", True), (" {{ ciclo }}. ", False), ("Grupo:", True), (" {{ grupo }}. ", False),
            ("Curso:", True), (" {{ curso_academico }}", False),
            ("\nCentro:", True), (" {{ centro_linea }}. ", False), ("Departamento:", True), (" {{ departamento_area }}", False),
        ]
        for texto, negrita in trozos:
            r = p.add_run(texto)
            r.bold = negrita
            r.font.name = "Arial"
    elif p.style.name == "Heading 1":
        for r in p.runs:
            r.font.size = Pt(9)
    elif t.startswith("{{ texto_") and "texto_final" not in t:
        p.paragraph_format.left_indent = Cm(0.5)
        for r in p.runs:
            r.bold = None

doc.save(RUTA)
print("ok", RUTA)
